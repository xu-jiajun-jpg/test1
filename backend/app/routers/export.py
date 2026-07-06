"""数据导出 — F11"""
import os
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from ..database import get_db
from ..middleware.auth_middleware import get_current_user_id
from ..models.sorting import SortingDecision
from ..models.stats import ExportRecord
from ..config import UPLOADS_DIR
import openpyxl
from fpdf import FPDF

router = APIRouter(prefix="/api/export", tags=["数据导出"])

@router.get("/excel")
def export_excel(page: int = 1, page_size: int = 20, user_id: int = Depends(get_current_user_id), db: Session = Depends(get_db)):
    if page_size > 10000: raise HTTPException(400, "请分批导出")
    decisions = db.query(SortingDecision).order_by(SortingDecision.decision_time.desc()).offset((page - 1) * page_size).limit(page_size).all()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "分拣记录"
    ws.append(["运单号", "分区", "城市", "分拣口", "状态", "决策时间"])
    for d in decisions: ws.append([d.tracking_number, d.district or d.province, d.city, d.target_chute, d.status, d.decision_time.isoformat() if d.decision_time else ""])
    path = UPLOADS_DIR / "export.xlsx"; wb.save(path)
    return FileResponse(path, filename="分拣记录.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@router.get("/pdf")
def export_pdf(page_size: int = 50, user_id: int = Depends(get_current_user_id), db: Session = Depends(get_db)):
    decisions = db.query(SortingDecision).order_by(SortingDecision.decision_time.desc()).limit(page_size).all()
    pdf = FPDF(); pdf.add_page()
    font_ok = False
    for fp in ["C:/Windows/Fonts/simhei.ttf", "C:/Windows/Fonts/simsun.ttc"]:
        if os.path.exists(fp):
            try:
                pdf.add_font("CJK", "", fp, uni=True)
                pdf.set_font("CJK", "", 12)
                font_ok = True
                break
            except:
                pass
    if not font_ok:
        pdf.set_font("Helvetica", "", 10)
    w = pdf.w - 2 * pdf.l_margin
    pdf.cell(w, 12, "物流分拣平台 - 分拣记录", ln=True, align="C")
    pdf.ln(4)
    for d in decisions:
        try:
            line = f"{d.tracking_number or '-'}  {d.district or d.province or '-'}  {d.target_chute or '-'}  {d.status or '-'}"
            pdf.cell(w, 8, line[:80], ln=True)
        except:
            pass
    path = UPLOADS_DIR / "export.pdf"
    pdf.output(str(path))
    return FileResponse(str(path), filename="分拣记录.pdf", media_type="application/pdf")
